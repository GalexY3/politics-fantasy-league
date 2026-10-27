"""Reproducible aggregate-only profile of a previously explored public XLSX.

Run with a previously downloaded original file and the bundled openpyxl runtime.
No names, identities, automatic matching or political scores are exported.
"""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import sys

from openpyxl import load_workbook


def main(source: Path, output: Path) -> None:
    blob = source.read_bytes()
    workbook = load_workbook(source, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise ValueError("Expected one sheet; inspect schema drift")
    sheet = workbook.active
    expected = ["Fecha", "Código elección", "Tipo elección", "Código Provincia", "Provincia",
                "Cargos Electos", "Candidatura", "Siglas Candidatura"]
    headers = list(next(sheet.iter_rows(min_row=4, max_row=4, values_only=True)))
    if headers[:len(expected)] != expected:
        raise ValueError(f"Unexpected column names: {headers}")
    elections, names, row_count, incomplete = Counter(), Counter(), 0, 0
    for row in sheet.iter_rows(min_row=5, values_only=True):
        if not any(value is not None for value in row):
            continue
        row_count += 1
        when, code, kind, province_code, province, elected, candidacy, acronym = row[:8]
        if any(value is None for value in (when, code, province_code, elected)):
            incomplete += 1
        key = when.isoformat() if isinstance(when, (datetime, date)) else str(when)
        elections[key] += 1
        if isinstance(elected, str) and elected.strip():
            names[elected.strip()] += 1
    result = {
        "id": "RC01-20261027", "source_url": "https://descargas.interior.gob.es/datasets/resultados_electorales/Elecciones-Congreso-cargos-electos.xlsx",
        "sha256": hashlib.sha256(blob).hexdigest(), "bytes": len(blob),
        "sheet_count": len(workbook.sheetnames), "sheet": sheet.title,
        "header_row": 4, "columns": headers, "data_rows": row_count,
        "rows_missing_core_fields": incomplete, "election_date_counts": dict(sorted(elections.items())),
        "exact_distinct_name_strings": len(names),
        "exact_name_strings_repeated_in_file": sum(n > 1 for n in names.values()),
        "semantics": "Election results per district; same name string is not a verified person or term of office.",
        "first_public_at_per_row": "NOT_PROVIDED_OR_UNVERIFIED", "identity_crosswalk_verified": False,
        "exploration_before_preregistration": True,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: python audit_interior_elected_20261027.py SOURCE.xlsx OUTPUT.json")
    main(Path(sys.argv[1]), Path(sys.argv[2]))
